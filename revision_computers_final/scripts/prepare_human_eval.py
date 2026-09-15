"""Build the blinded 320-text, three-rater human-evaluation package."""

from __future__ import annotations

import csv
import hashlib
import json
import random
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from common import REVISION_ROOT, atomic_write_csv, atomic_write_text, read_jsonl_tolerant  # type: ignore
else:
    from .common import REVISION_ROOT, atomic_write_csv, atomic_write_text, read_jsonl_tolerant


REPRESENTATIVE_MODELS = ["qwen35_4b", "qwen35_9b", "qwen36_27b", "qwen36_35b_a3b"]
INTERNAL_RATIOS = [1 / 6, 2 / 6, 3 / 6, 4 / 6, 5 / 6]
SELECTION_SEED = 20260916


def stable_digest(*parts: object) -> str:
    return hashlib.sha256("::".join(map(str, parts)).encode("utf-8")).hexdigest()


def load_generations() -> dict[tuple[str, str, float, int], dict[str, Any]]:
    output: dict[tuple[str, str, float, int], dict[str, Any]] = {}
    for model_id in REPRESENTATIVE_MODELS:
        path = REVISION_ROOT / "data/fragments/main" / f"{model_id}.jsonl"
        source = read_jsonl_tolerant(path)
        if source.bad_lines or len(source.rows) != 6720:
            raise ValueError(f"main output incomplete for human package: {model_id}")
        for row in source.rows:
            key = (model_id, row["scenario_id"], float(row["target_ratio"]), int(row["seed"]))
            output[key] = row
    return output


def main() -> None:
    with (REVISION_ROOT / "config/human_eval_scenarios.csv").open("r", encoding="utf-8", newline="") as handle:
        scenarios = list(csv.DictReader(handle))
    if len(scenarios) != 16:
        raise ValueError("human selection must contain 16 scenarios")
    generations = load_generations()
    items: list[dict[str, Any]] = []
    private: list[dict[str, Any]] = []
    for scenario in scenarios:
        for model_id in REPRESENTATIVE_MODELS:
            for ratio in INTERNAL_RATIOS:
                seed = 1 + int(stable_digest(SELECTION_SEED, scenario["scenario_id"], model_id, ratio)[:8], 16) % 3
                row = generations[(model_id, scenario["scenario_id"], ratio, seed)]
                item_id = "HE_" + stable_digest("human_eval", row["run_id"])[:12].upper()
                items.append(
                    {
                        "item_id": item_id,
                        "scenario_id": scenario["scenario_id"],
                        "style_axis": scenario["style_axis"],
                        "input_text": scenario["input_text"],
                        "style_A": scenario["style_A"],
                        "style_B": scenario["style_B"],
                        "text_to_rate": row["generated_text"],
                        "attention_check": "no",
                    }
                )
                private.append(
                    {
                        "item_id": item_id,
                        "run_id": row["run_id"],
                        "scenario_id": scenario["scenario_id"],
                        "style_axis": scenario["style_axis"],
                        "model_id": model_id,
                        "target_ratio": ratio,
                        "seed": seed,
                        "reused_legacy_item": "no",
                    }
                )
    if len(items) != 320 or len({row["item_id"] for row in items}) != 320:
        raise AssertionError("human package must contain 320 unique texts")
    randomizer = random.Random(SELECTION_SEED)
    randomizer.shuffle(items)
    for index, row in enumerate(items, 1):
        row["global_display_order"] = index
    order_by_id = {row["item_id"]: row["global_display_order"] for row in items}
    for row in private:
        row["global_display_order"] = order_by_id[row["item_id"]]

    public_fields = [
        "item_id", "global_display_order", "scenario_id", "style_axis", "input_text",
        "style_A", "style_B", "text_to_rate", "attention_check",
    ]
    private_fields = [
        "item_id", "global_display_order", "run_id", "scenario_id", "style_axis",
        "model_id", "target_ratio", "seed", "reused_legacy_item",
    ]
    out_dir = REVISION_ROOT / "human_eval"
    atomic_write_csv(out_dir / "annotation_items.csv", public_fields, items)
    atomic_write_csv(out_dir / "randomization_map.csv", private_fields, sorted(private, key=lambda row: row["global_display_order"]))

    result_fields = [
        "assignment_id", "item_id", "rater_slot", "anonymous_rater_id",
        "style_B_intensity_0_100", "content_coverage_1_5", "naturalness_1_5",
        "adds_key_information_yes_no", "contradicts_source_yes_no", "unable_to_judge_yes_no",
        "start_time_utc", "end_time_utc", "duration_seconds", "rater_comment",
        "attention_check", "attention_check_expected",
    ]
    assignments: list[dict[str, Any]] = []
    item_by_id = {row["item_id"]: row for row in items}
    for slot in ("A", "B", "C"):
        ids = list(item_by_id)
        random.Random(f"{SELECTION_SEED}:{slot}").shuffle(ids)
        for position, item_id in enumerate(ids, 1):
            assignments.append(
                {
                    "assignment_id": f"{slot}_{position:03d}",
                    "item_id": item_id,
                    "rater_slot": slot,
                    **{field: "" for field in result_fields[3:14]},
                    "attention_check": "no",
                    "attention_check_expected": "",
                }
            )
        for check_index in range(1, 9):
            assignments.append(
                {
                    "assignment_id": f"{slot}_AC{check_index:02d}",
                    "item_id": f"ATTENTION_{slot}_{check_index:02d}",
                    "rater_slot": slot,
                    **{field: "" for field in result_fields[3:14]},
                    "attention_check": "yes",
                    "attention_check_expected": "unable_to_judge_yes_no=yes",
                }
            )
    atomic_write_csv(out_dir / "human_results.csv", result_fields, assignments)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for slot in ("A", "B", "C"):
        sheet = workbook.create_sheet(f"rater_{slot}")
        headers = [
            "assignment_id", "item_id", "anonymous_rater_id", "input_text", "style_A",
            "style_B", "text_to_rate", "style_B_intensity_0_100", "content_coverage_1_5",
            "naturalness_1_5", "adds_key_information_yes_no", "contradicts_source_yes_no",
            "unable_to_judge_yes_no", "start_time_utc", "end_time_utc", "duration_seconds",
            "rater_comment", "attention_check",
        ]
        sheet.append(headers)
        slot_assignments = [row for row in assignments if row["rater_slot"] == slot]
        random.Random(f"form:{SELECTION_SEED}:{slot}").shuffle(slot_assignments)
        for assignment in slot_assignments:
            if assignment["attention_check"] == "yes":
                item = {
                    "input_text": "注意力检查：本题不评价正文。",
                    "style_A": "注意力检查",
                    "style_B": "注意力检查",
                    "text_to_rate": "请在“无法判断”一栏选择 yes，其余评分保持空白。",
                }
            else:
                item = item_by_id[assignment["item_id"]]
            sheet.append(
                [
                    assignment["assignment_id"], assignment["item_id"], "",
                    item["input_text"], item["style_A"], item["style_B"], item["text_to_rate"],
                    "", "", "", "", "", "", "", "", "", "", assignment["attention_check"],
                ]
            )
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        for column in (4, 5, 6, 7, 17):
            sheet.column_dimensions[get_column_letter(column)].width = 45
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        intensity = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
        score = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
        yes_no = DataValidation(type="list", formula1='"yes,no"')
        sheet.add_data_validation(intensity)
        sheet.add_data_validation(score)
        sheet.add_data_validation(yes_no)
        intensity.add(f"H2:H{sheet.max_row}")
        score.add(f"I2:J{sheet.max_row}")
        for column in ("K", "L", "M"):
            yes_no.add(f"{column}2:{column}{sheet.max_row}")
    workbook.save(out_dir / "annotation_form.xlsx")

    instructions = """# 连续风格强度人工评价说明

## 任务与盲法

每位评价者完成一个 `rater_A/B/C` 工作表。模型、目标比例和生成 seed 均已隐藏，文本顺序随机。请勿尝试识别模型。每条正式文本共有三份独立评价；320 条唯一文本与 960 次正式评分任务分别计数，不把重复评价当作新文本。

## 每条文本的评分

1. `style_B_intensity_0_100`：相对给定风格 A/B，判断文本呈现风格 B 的强度。0 表示完全接近 A，100 表示完全接近 B。
2. `content_coverage_1_5`：固定核心语义覆盖程度，1 为严重缺失，5 为完整覆盖。
3. `naturalness_1_5`：中文表达自然度，1 为明显不自然，5 为自然流畅。
4. `adds_key_information_yes_no`：是否加入了会改变任务结果或核心事实的关键信息。普通解释性细节不算。
5. `contradicts_source_yes_no`：是否与固定核心语义发生事实或任务目标矛盾。
6. `unable_to_judge_yes_no`：信息不足、文本损坏或无法可靠判断时选 yes，并将其他评分留空。

填写匿名 rater_id、开始和结束 UTC 时间；`duration_seconds` 为两者之差。不要在评论中写姓名或联系方式。遇到注意力检查时按题面操作。

## 预先固定的纳入与排除规则

- 纳入：成年、中文熟练、阅读并同意匿名研究说明、独立完成分配任务。
- 正式条目缺任一必填项时，该条评价无效；选择“无法判断”时其他量表应留空。
- 单条正式评价用时少于 5 秒标为速度异常；不自动删除，而在主分析和排除该条的敏感性分析中分别报告。
- 每个工作表含 8 个注意力检查。失败 2 个或以上时排除该评价者整批正式评价；失败 1 个保留并作敏感性分析。
- 同一匿名 rater_id 不得填写同一 item_id 两次。无法确认独立作答或存在批量复制时整批排除并记录理由。
- 排除决定必须在揭盲模型和比例前完成并留日志。

## 数据安全

只保存匿名 rater_id、评分与时长，不收集姓名、邮箱或设备标识。`randomization_map.csv` 含模型与比例映射，不发给评价者；评分冻结后再用于统计合并。
"""
    atomic_write_text(out_dir / "instructions_zh.md", instructions)
    print(json.dumps({"unique_texts": 320, "formal_assignments": 960, "attention_checks": 24, "total_form_rows": 984}, ensure_ascii=False))


if __name__ == "__main__":
    main()
