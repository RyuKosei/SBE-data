"""Create the blinded two-author shared-anchor review package.

Reference text cells intentionally remain empty: fabricating author-reviewed
anchors would violate the protocol. The workbook is immediately usable for
independent human rewriting and review, while the CSV remains explicitly
``external_pending`` until both checks pass.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from common import REVISION_ROOT, atomic_write_csv, read_jsonl_tolerant  # type: ignore
else:
    from .common import REVISION_ROOT, atomic_write_csv, read_jsonl_tolerant


SCENARIO_FILE = Path("/home/data_cpfs/lizihua/sbe/data/scenarios/styleblend_8axis_scenarios_v1.jsonl")


def main() -> None:
    selection_path = REVISION_ROOT / "config/shared_anchor_scenarios.csv"
    with selection_path.open("r", encoding="utf-8", newline="") as handle:
        selected = list(csv.DictReader(handle))
    source = read_jsonl_tolerant(SCENARIO_FILE)
    source_by_id = {row["scenario_id"]: row for row in source.rows}
    if len(selected) != 80:
        raise ValueError("shared-anchor package requires exactly 80 scenarios")

    fields = [
        "scenario_id", "style_axis", "input_text", "style_A", "style_B",
        "content_checks", "reference_text_A", "reference_text_B",
        "author_1_content_consistent", "author_1_style_isolated",
        "author_1_separation_sufficient", "author_1_length_not_dominant",
        "author_1_no_new_fact_or_contradiction", "author_1_notes",
        "author_2_content_consistent", "author_2_style_isolated",
        "author_2_separation_sufficient", "author_2_length_not_dominant",
        "author_2_no_new_fact_or_contradiction", "author_2_notes",
        "approval_status",
    ]
    rows = []
    for selected_row in selected:
        scenario = source_by_id[selected_row["scenario_id"]]
        rows.append(
            {
                "scenario_id": scenario["scenario_id"],
                "style_axis": scenario["axis_id"],
                "input_text": scenario["content"],
                "style_A": scenario["style_a"],
                "style_B": scenario["style_b"],
                "content_checks": "；".join(scenario.get("content_checks", [])),
                "reference_text_A": "",
                "reference_text_B": "",
                **{field: "" for field in fields if field.startswith("author_")},
                "approval_status": "external_pending",
            }
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "shared_anchor_candidates"
    sheet.append(fields)
    for row in rows:
        sheet.append([row.get(field, "") for field in fields])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 24, "B": 20, "C": 38, "D": 34, "E": 34, "F": 30,
        "G": 55, "H": 55, "N": 34, "T": 34, "U": 20,
    }
    for column in range(1, len(fields) + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = widths.get(letter, 18)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    validation = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True)
    sheet.add_data_validation(validation)
    for column in (9, 10, 11, 12, 13, 15, 16, 17, 18, 19):
        validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}81")

    guide = workbook.create_sheet("instructions")
    instructions = [
        "Shared-anchor author workflow",
        "1. One author independently writes A and B endpoint texts from input_text; do not copy any tested model output.",
        "2. Both texts must preserve every content checkpoint and differ mainly on the named style axis.",
        "3. A second author checks factual consistency, style isolation, endpoint separation, length dominance, and new facts/contradictions.",
        "4. Author identities and dates should be kept in the private study log, not in the public text file.",
        "5. approval_status may become approved only when both authors mark all five checks yes and both reference texts are nonempty.",
        "6. Run the validator before producing shared_anchors_final.csv. No unchecked row may enter model comparison.",
    ]
    for line in instructions:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=14)
    for cell in guide["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook_path = REVISION_ROOT / "data/shared_anchor_candidates.xlsx"
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    atomic_write_csv(REVISION_ROOT / "data/shared_anchors_final.csv", fields, rows)
    print(f"wrote {len(rows)} scenario pairs / {len(rows) * 2} endpoint slots; all remain external_pending")


if __name__ == "__main__":
    main()
